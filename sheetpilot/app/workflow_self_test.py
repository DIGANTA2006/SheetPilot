"""Synthetic, local-only release self-tests for the deterministic workflow boundary."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

import openpyxl

from sheetpilot.app.config import AppConfig
from sheetpilot.core.audit_report import AuditReport
from sheetpilot.core.backup_service import BackupService
from sheetpilot.core.exceptions import SheetPilotError, UnknownOperationError
from sheetpilot.core.executor import JobExecutor
from sheetpilot.core.plan_schema import (
    OperationPlan,
    OutputFormat,
    OutputSettings,
    PlanStep,
    SourceReference,
    StepTarget,
)
from sheetpilot.core.preview_engine import ExecutionApproval, PreviewEngine, SourceBinding
from sheetpilot.core.reconciliation import ReconciliationStatus
from sheetpilot.engines.tabular_io import validate_output_file
from sheetpilot.operations.registry import build_default_registry
from sheetpilot.security.hashing import fingerprint_file, verify_fingerprint


class WorkflowSelfTestError(SheetPilotError):
    """A release self-test failed to prove a required safety property."""

    code = "workflow_self_test_failed"


@dataclass(frozen=True)
class NormalWorkflowSelfTestEvidence:
    """Non-sensitive aggregate evidence from the synthetic normal workflow."""

    source_sha256: str
    backup_sha256: str
    output_sha256: str
    audit_sha256: str
    original_row_count: int
    final_row_count: int
    rows_changed: int
    reconciliation_status: str


@dataclass(frozen=True)
class InvalidWorkflowSelfTestEvidence:
    """Evidence that an unregistered operation failed before artifact creation."""

    source_sha256: str
    rejected_operation: str
    artifact_count: int


def _require(condition: bool, message: str) -> None:
    if not condition:
        raise WorkflowSelfTestError(message)


def _config(root: Path) -> AppConfig:
    data_root = root / "app-data"
    return AppConfig(
        data_dir=data_root,
        backup_dir=root / "backups",
        temp_dir=root / "temporary-workspaces",
        database_path=data_root / "sheetpilot.sqlite3",
        log_dir=root / "logs",
    )


def _create_synthetic_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    if worksheet is None:
        workbook.close()
        raise WorkflowSelfTestError("The synthetic workbook has no active worksheet.")
    worksheet.title = "Customers"
    worksheet.append(["ID", "Name"])
    worksheet.append([1, " Ada Lovelace "])
    worksheet.append([2, " Grace Hopper "])
    workbook.save(path)
    workbook.close()


def _normal_plan(source: Path) -> tuple[OperationPlan, SourceBinding]:
    source_id = uuid4()
    fingerprint = fingerprint_file(source)
    plan = OperationPlan(
        job_name="SheetPilot synthetic XLSX self-test",
        source_files=[
            SourceReference(
                source_id=source_id,
                file_name=source.name,
                sha256=fingerprint.sha256,
                sheet_names=["Customers"],
            )
        ],
        steps=[
            PlanStep(
                step_id="trim-customer-names",
                operation="text.clean",
                parameters={"columns": ["Name"], "actions": [{"kind": "trim"}]},
                target=StepTarget(
                    source_id=source_id,
                    sheet="Customers",
                    columns=["Name"],
                ),
                explanation="Trim whitespace in synthetic customer names.",
            )
        ],
        output=OutputSettings(
            output_name="sheetpilot-self-test-result",
            format=OutputFormat.XLSX,
            preserve_formatting=True,
        ),
    )
    return plan, SourceBinding(source_id=source_id, path=source, fingerprint=fingerprint)


def _verify_reopened_output(path: Path) -> None:
    validate_output_file(path)
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        _require("Customers" in workbook.sheetnames, "The output lost the synthetic worksheet.")
        worksheet = workbook["Customers"]
        _require(worksheet["A1"].value == "ID", "The output header was not preserved.")
        _require(worksheet["B1"].value == "Name", "The output header was not preserved.")
        _require(worksheet["B2"].value == "Ada Lovelace", "The first trim was not exported.")
        _require(worksheet["B3"].value == "Grace Hopper", "The second trim was not exported.")
    finally:
        workbook.close()


def _exercise_normal_workflow(root: Path) -> NormalWorkflowSelfTestEvidence:
    source = root / "synthetic-source.xlsx"
    output_directory = root / "output"
    _create_synthetic_workbook(source)
    source_bytes = source.read_bytes()
    source_mtime_ns = source.stat().st_mtime_ns
    plan, binding = _normal_plan(source)
    registry = build_default_registry()
    preview, _ = PreviewEngine(registry).generate(plan, (binding,))
    _require(preview.total_change_count == 2, "The normal preview did not find both changes.")

    approval = ExecutionApproval(
        job_id=plan.job_id,
        preview_digest=preview.preview_digest,
        approved=True,
    )
    result = JobExecutor(_config(root), registry).execute(
        plan,
        (binding,),
        preview,
        approval,
        output_directory,
    )

    _require(source.read_bytes() == source_bytes, "The normal self-test modified its source.")
    _require(
        source.stat().st_mtime_ns == source_mtime_ns,
        "The normal self-test changed its source timestamp.",
    )
    verify_fingerprint(source, binding.fingerprint)
    _require(len(result.backups) == 1, "The normal self-test did not create one backup.")
    backup = result.backups[0]
    _require(backup.backup_path.is_file(), "The verified backup is missing.")
    _require(backup.manifest_path.is_file(), "The backup manifest is missing.")
    verify_fingerprint(backup.backup_path, binding.fingerprint)
    _require(backup.backup_path.read_bytes() == source_bytes, "The backup differs from the source.")
    loaded_backup = BackupService(_config(root).backup_dir).load_receipt(backup.manifest_path)
    _require(loaded_backup == backup, "The backup manifest did not reopen to the same receipt.")

    _require(result.output.path.is_file(), "The normal output is missing.")
    _require(result.output.path != source, "The normal output aliases the source.")
    verify_fingerprint(result.output.path, result.output.fingerprint)
    _verify_reopened_output(result.output.path)

    _require(result.audit.path.is_file(), "The normal audit report is missing.")
    verify_fingerprint(result.audit.path, result.audit.fingerprint)
    audit = AuditReport.model_validate_json(result.audit.path.read_text(encoding="utf-8"))
    _require(audit.job_id == str(plan.job_id), "The audit report belongs to another job.")
    _require(
        audit.output.sha256 == result.output.fingerprint.sha256,
        "The audit output hash does not match the committed output.",
    )
    _require(
        audit.sources[0].sha256 == binding.fingerprint.sha256,
        "The audit source hash does not match the analysed source.",
    )

    reconciliation = result.reconciliation
    _require(reconciliation.original_row_count == 2, "The original row count did not reconcile.")
    _require(reconciliation.final_row_count == 2, "The final row count did not reconcile.")
    _require(reconciliation.rows_changed == 2, "The changed-row count did not reconcile.")
    _require(
        reconciliation.status == ReconciliationStatus.PASSED,
        "The normal workflow reconciliation did not pass.",
    )
    _require(
        audit.reconciliation == reconciliation,
        "The audit reconciliation differs from the execution result.",
    )
    return NormalWorkflowSelfTestEvidence(
        source_sha256=binding.fingerprint.sha256,
        backup_sha256=backup.source_fingerprint.sha256,
        output_sha256=result.output.fingerprint.sha256,
        audit_sha256=result.audit.fingerprint.sha256,
        original_row_count=reconciliation.original_row_count,
        final_row_count=reconciliation.final_row_count,
        rows_changed=reconciliation.rows_changed,
        reconciliation_status=reconciliation.status.value,
    )


def run_normal_workflow_self_test() -> NormalWorkflowSelfTestEvidence:
    """Run a complete XLSX transaction using only disposable synthetic data."""
    with TemporaryDirectory(prefix="sheetpilot-normal-self-test-") as directory:
        return _exercise_normal_workflow(Path(directory))


def _exercise_invalid_workflow(root: Path) -> InvalidWorkflowSelfTestEvidence:
    source = root / "synthetic-invalid-source.xlsx"
    _create_synthetic_workbook(source)
    source_bytes = source.read_bytes()
    source_mtime_ns = source.stat().st_mtime_ns
    source_id = uuid4()
    fingerprint = fingerprint_file(source)
    rejected_operation = "selftest.unknown-operation"
    plan = OperationPlan(
        job_name="SheetPilot intentionally invalid self-test",
        source_files=[
            SourceReference(
                source_id=source_id,
                file_name=source.name,
                sha256=fingerprint.sha256,
                sheet_names=["Customers"],
            )
        ],
        steps=[
            PlanStep(
                step_id="reject-unknown-operation",
                operation=rejected_operation,
                target=StepTarget(source_id=source_id, sheet="Customers"),
                explanation="This unregistered operation must be rejected.",
            )
        ],
        output=OutputSettings(
            output_name="invalid-workflow-must-not-exist",
            format=OutputFormat.XLSX,
            preserve_formatting=True,
        ),
    )
    binding = SourceBinding(source_id=source_id, path=source, fingerprint=fingerprint)
    try:
        PreviewEngine(build_default_registry()).generate(plan, (binding,))
    except UnknownOperationError as error:
        _require(
            rejected_operation in str(error),
            "The invalid workflow rejected an unexpected operation.",
        )
    else:
        raise WorkflowSelfTestError("The registry accepted an unregistered operation.")

    _require(source.read_bytes() == source_bytes, "The invalid self-test modified its source.")
    _require(
        source.stat().st_mtime_ns == source_mtime_ns,
        "The invalid self-test changed its source timestamp.",
    )
    verify_fingerprint(source, fingerprint)
    artifacts = tuple(path for path in root.rglob("*") if path.is_file() and path != source)
    _require(not artifacts, "The rejected workflow created an output, backup, or audit artifact.")
    return InvalidWorkflowSelfTestEvidence(
        source_sha256=fingerprint.sha256,
        rejected_operation=rejected_operation,
        artifact_count=len(artifacts),
    )


def run_invalid_workflow_self_test() -> InvalidWorkflowSelfTestEvidence:
    """Prove an unknown operation is rejected without creating durable artifacts."""
    with TemporaryDirectory(prefix="sheetpilot-invalid-self-test-") as directory:
        return _exercise_invalid_workflow(Path(directory))
