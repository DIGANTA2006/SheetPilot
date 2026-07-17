"""Bounded tabular input and output validation adapters."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import polars as pl

from sheetpilot.app.config import SecurityLimits
from sheetpilot.core.exceptions import (
    CorruptWorkbookError,
    InvalidPlanError,
    UnsupportedFormatError,
)
from sheetpilot.engines.csv_engine import read_csv
from sheetpilot.engines.openpyxl_safety import close_workbook
from sheetpilot.security.archive_guard import inspect_ooxml_archive


def read_workbook_table(path: Path, sheet_name: str) -> pl.DataFrame:
    """Read a worksheet table without evaluating formulas or links."""
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=False,
        keep_vba=path.suffix.casefold() == ".xlsm",
        keep_links=False,
    )
    try:
        if sheet_name not in workbook.sheetnames:
            raise InvalidPlanError(f"Worksheet is missing: {sheet_name}")
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
    finally:
        close_workbook(workbook)
    if not rows:
        return pl.DataFrame()
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    if not all(headers) or len(headers) != len(set(header.casefold() for header in headers)):
        raise InvalidPlanError("Worksheet tables require non-blank unique headers.")
    values = [row for row in rows[1:] if any(value is not None for value in row)]
    return pl.DataFrame(values, schema=headers, orient="row", strict=False)


def read_tabular_source(path: Path, *, sheet_name: str | None = None) -> pl.DataFrame:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        if sheet_name is None:
            workbook = openpyxl.load_workbook(path, read_only=True, keep_links=False)
            try:
                if not workbook.sheetnames:
                    raise CorruptWorkbookError("The workbook has no worksheets.")
                sheet_name = workbook.sheetnames[0]
            finally:
                workbook.close()
        return read_workbook_table(path, sheet_name)
    raise UnsupportedFormatError("Supported tabular inputs are .csv, .xlsx, and .xlsm.")


def validate_output_file(path: Path, limits: SecurityLimits | None = None) -> None:
    """Reopen an output through its real parser before publication."""
    if path.suffix.casefold() == ".csv":
        read_csv(path)
        return
    if path.suffix.casefold() in {".xlsx", ".xlsm"}:
        inspect_ooxml_archive(path, limits or SecurityLimits())
        workbook = openpyxl.load_workbook(path, read_only=True, keep_links=False)
        try:
            if not workbook.sheetnames:
                raise CorruptWorkbookError("The generated workbook has no worksheets.")
        finally:
            workbook.close()
        return
    raise UnsupportedFormatError("The generated output format is unsupported.")
