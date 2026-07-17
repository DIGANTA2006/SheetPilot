"""Copy-only existing-workbook edits with style and VBA preservation."""

from __future__ import annotations

import shutil
from copy import copy
from enum import StrEnum
from pathlib import Path
from typing import Any, cast

import openpyxl
import polars as pl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import BaseModel, ConfigDict

from sheetpilot.core.exceptions import InvalidPlanError, OutputCollisionError
from sheetpilot.engines.formulas import GeneratedFormulaSpec, render_formula
from sheetpilot.engines.openpyxl_safety import close_workbook
from sheetpilot.security.formula_guard import is_formula_injection


class NumberFormat(StrEnum):
    INTEGER = "0"
    DECIMAL_2 = "0.00"
    PERCENT_2 = "0.00%"
    CURRENCY_INR = "₹#,##0.00"
    DATE_ISO = "yyyy-mm-dd"
    DATE_INDIA = "dd-mm-yyyy"
    DATETIME_ISO = "yyyy-mm-dd hh:mm:ss"


class ColumnFormatSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    column: str
    number_format: NumberFormat


def _copy_style(source_cell: Cell, target_cell: Cell) -> None:
    source_any = cast(Any, source_cell)
    target_any = cast(Any, target_cell)
    target_any.font = copy(source_any.font)
    target_any.fill = copy(source_any.fill)
    target_any.border = copy(source_any.border)
    target_any.alignment = copy(source_any.alignment)
    target_any.number_format = source_any.number_format
    target_any.protection = copy(source_any.protection)


def _write_safe_cell(worksheet: Worksheet, row: int, column: int, value: Any) -> None:
    cell = worksheet.cell(row=row, column=column)
    cell.value = value
    if isinstance(value, str) and is_formula_injection(value):
        cell.data_type = "s"


def replace_sheet_table(
    worksheet: Worksheet,
    frame: pl.DataFrame,
    *,
    formula_specs: list[GeneratedFormulaSpec] | None = None,
    formats: list[ColumnFormatSpec] | None = None,
) -> None:
    """Replace cell values while retaining existing sheet structure and styles."""
    headers = list(frame.columns)
    formulas = formula_specs or []
    output_headers = headers + [spec.output_column for spec in formulas]
    if len(output_headers) != len(set(header.casefold() for header in output_headers)):
        raise InvalidPlanError("Workbook output requires unique headers.")
    max_rows = max(worksheet.max_row, frame.height + 1)
    max_columns = max(worksheet.max_column, len(output_headers))
    for row in worksheet.iter_rows(min_row=1, max_row=max_rows, max_col=max_columns):
        for cell in row:
            cell.value = None
    for column_index, header in enumerate(output_headers, start=1):
        if column_index > 1 and worksheet.cell(1, 1).has_style:
            source_header = worksheet.cell(1, 1)
            target_header = worksheet.cell(1, column_index)
            if isinstance(source_header, Cell) and isinstance(target_header, Cell):
                _copy_style(source_header, target_header)
        _write_safe_cell(worksheet, 1, column_index, header)
    for row_index, row in enumerate(frame.iter_rows(), start=2):
        for column_index, value in enumerate(row, start=1):
            _write_safe_cell(worksheet, row_index, column_index, value)
        for formula_index, spec in enumerate(formulas, start=len(headers) + 1):
            cell = worksheet.cell(row=row_index, column=formula_index)
            cell.value = render_formula(spec, headers, row_index)
            if spec.number_format:
                cell.number_format = spec.number_format
    positions = {header: index + 1 for index, header in enumerate(output_headers)}
    for format_spec in formats or []:
        if format_spec.column not in positions:
            raise InvalidPlanError(f"Format column is missing: {format_spec.column}")
        for row_index in range(2, frame.height + 2):
            worksheet.cell(
                row=row_index, column=positions[format_spec.column]
            ).number_format = format_spec.number_format.value
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        f"A1:{openpyxl.utils.get_column_letter(len(output_headers))}{frame.height + 1}"
    )


def modify_workbook_copy(
    source_copy: Path,
    destination: Path,
    *,
    replacements: dict[str, pl.DataFrame],
    sheet_renames: dict[str, str] | None = None,
    formula_specs: dict[str, list[GeneratedFormulaSpec]] | None = None,
    formats: dict[str, list[ColumnFormatSpec]] | None = None,
) -> Path:
    """Modify only a new copy; the supplied source copy is opened read-only in intent."""
    if destination.exists():
        raise OutputCollisionError("Workbook modification cannot overwrite an output.")
    if source_copy.resolve() == destination.resolve():
        raise InvalidPlanError("Source and destination must be distinct.")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_copy, destination)
    keep_vba = source_copy.suffix.casefold() == ".xlsm"
    workbook = openpyxl.load_workbook(destination, keep_vba=keep_vba, keep_links=False)
    try:
        renames = sheet_renames or {}
        missing_renames = sorted(set(renames) - set(workbook.sheetnames))
        if missing_renames:
            raise InvalidPlanError(f"Sheets to rename are missing: {', '.join(missing_renames)}")
        final_names = [renames.get(name, name) for name in workbook.sheetnames]
        if len(final_names) != len(set(name.casefold() for name in final_names)):
            raise InvalidPlanError("Sheet renames would create duplicate names.")
        for old, new in renames.items():
            workbook[old].title = new
        for requested_name, frame in replacements.items():
            sheet_name = renames.get(requested_name, requested_name)
            worksheet = (
                workbook[sheet_name]
                if sheet_name in workbook.sheetnames
                else workbook.create_sheet(sheet_name)
            )
            replace_sheet_table(
                worksheet,
                frame,
                formula_specs=(formula_specs or {}).get(requested_name),
                formats=(formats or {}).get(requested_name),
            )
        workbook.save(destination)
    except BaseException:
        close_workbook(workbook)
        destination.unlink(missing_ok=True)
        raise
    close_workbook(workbook)
    return destination
