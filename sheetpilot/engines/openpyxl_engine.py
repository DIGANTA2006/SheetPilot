"""Bounded, read-only-behaviour Excel workbook profiling."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from sheetpilot.app.config import SecurityLimits
from sheetpilot.core.exceptions import CorruptWorkbookError, FileLimitError
from sheetpilot.core.profile_models import SheetProfile
from sheetpilot.engines.profile_helpers import profile_rows


@dataclass(frozen=True)
class WorkbookAnalysis:
    sheets: tuple[SheetProfile, ...]
    named_ranges: tuple[str, ...]
    workbook_protected: bool
    external_formula_present: bool
    estimated_memory_bytes: int


def _trim_used_rows(rows: list[tuple[Any, ...]]) -> tuple[list[tuple[Any, ...]], int]:
    last_row = 0
    last_column = 0
    for row_index, row in enumerate(rows, start=1):
        row_has_value = False
        for column_index, value in enumerate(row, start=1):
            if value is not None and (not isinstance(value, str) or value.strip()):
                row_has_value = True
                last_column = max(last_column, column_index)
        if row_has_value:
            last_row = row_index
    return ([row[:last_column] for row in rows[:last_row]], last_column)


def profile_workbook(path: Path, limits: SecurityLimits) -> WorkbookAnalysis:
    """Profile workbook structure without saving, recalculating, or refreshing links."""
    keep_vba = path.suffix.casefold() == ".xlsm"
    try:
        workbook = openpyxl.load_workbook(
            path,
            read_only=False,
            data_only=False,
            keep_vba=keep_vba,
            keep_links=False,
        )
    except (InvalidFileException, OSError, KeyError, ValueError, TypeError) as error:
        raise CorruptWorkbookError("The workbook could not be parsed safely.") from error
    try:
        estimated_total = 0
        sheet_profiles: list[SheetProfile] = []
        external_formula_present = False
        for worksheet in workbook.worksheets:
            if worksheet.max_row > limits.max_rows or worksheet.max_column > limits.max_columns:
                raise FileLimitError("A worksheet exceeds the configured row or column limit.")
            estimate = worksheet.max_row * worksheet.max_column * 120
            estimated_total += estimate
            if estimated_total > limits.max_estimated_memory_bytes:
                raise FileLimitError("The workbook exceeds the configured memory estimate limit.")
            cells = list(worksheet.iter_rows())
            values = [tuple(cell.value for cell in row) for row in cells]
            trimmed, used_columns = _trim_used_rows(values)
            used_rows = len(trimmed)
            formula_cells = 0
            formula_error_cells = 0
            for row in cells[:used_rows]:
                for cell in row[:used_columns]:
                    if cell.data_type == "f":
                        formula_cells += 1
                        formula = str(cell.value)
                        if "[" in formula and "]" in formula:
                            external_formula_present = True
                    elif cell.data_type == "e":
                        formula_error_cells += 1
            raw_headers = list(trimmed[0]) if trimmed else []
            data_rows = trimmed[1:] if trimmed else []
            sheet_profiles.append(
                profile_rows(
                    name=worksheet.title,
                    visibility=worksheet.sheet_state,
                    raw_headers=raw_headers,
                    rows=data_rows,
                    used_rows=used_rows,
                    used_columns=used_columns,
                    formula_cells=formula_cells,
                    formula_error_cells=formula_error_cells,
                    merged_ranges=tuple(str(item) for item in worksheet.merged_cells.ranges),
                    protected=bool(worksheet.protection.sheet),
                )
            )
        named_ranges = tuple(sorted(item.name for item in workbook.defined_names.values()))
        security = workbook.security
        workbook_protected = bool(security.lockStructure or security.lockWindows)
        return WorkbookAnalysis(
            sheets=tuple(sheet_profiles),
            named_ranges=named_ranges,
            workbook_protected=workbook_protected,
            external_formula_present=external_formula_present,
            estimated_memory_bytes=estimated_total,
        )
    finally:
        workbook.close()
