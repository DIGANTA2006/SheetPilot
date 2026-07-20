"""Reproducible before/after preview generation and approval binding."""

from __future__ import annotations

import hashlib
import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from uuid import UUID

import openpyxl
import polars as pl
from pydantic import BaseModel, ConfigDict, Field

from sheetpilot.core.change_models import ChangeKind, ChangeRecord
from sheetpilot.core.exceptions import InvalidPlanError, SourceChangedError
from sheetpilot.core.operation_registry import OperationRegistry
from sheetpilot.core.plan_runner import (
    ROW_ID_COLUMN,
    DatasetKey,
    PlanRunner,
    PlanRunResult,
    strip_internal_columns,
)
from sheetpilot.core.plan_schema import OperationPlan
from sheetpilot.engines.tabular_io import read_tabular_source, read_workbook_table
from sheetpilot.security.hashing import FileFingerprint, verify_fingerprint


class SourceBinding(BaseModel):
    """Bind a plan source ID to an analysed local file and exact fingerprint."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    source_id: UUID
    path: Path
    fingerprint: FileFingerprint


class PreviewStepSummary(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    step_id: str
    operation: str
    change_count: int = Field(ge=0)
    warnings: tuple[str, ...]


class PreviewResult(BaseModel):
    """Approval-safe preview metadata plus an in-memory sample of change values."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    job_id: UUID
    created_at: datetime
    changes: tuple[ChangeRecord, ...]
    total_change_count: int = Field(ge=0)
    sampled: bool
    steps: tuple[PreviewStepSummary, ...]
    plan_digest: str = Field(pattern=r"^[a-f0-9]{64}$")
    source_digest: str = Field(pattern=r"^[a-f0-9]{64}$")
    result_digest: str = Field(pattern=r"^[a-f0-9]{64}$")
    preview_digest: str = Field(pattern=r"^[a-f0-9]{64}$")


class ExecutionApproval(BaseModel):
    """Explicit human decision cryptographically bound to one preview."""

    model_config = ConfigDict(extra="forbid", frozen=True)

    job_id: UUID
    preview_digest: str = Field(pattern=r"^[a-f0-9]{64}$")
    approved: bool
    confirmed_destructive_steps: frozenset[str] = frozenset()
    rejected_change_ids: frozenset[str] = frozenset()
    approved_at: datetime = Field(default_factory=lambda: datetime.now(UTC))


def _json_default(value: Any) -> str:
    if isinstance(value, (datetime, Path, UUID)):
        return str(value)
    if hasattr(value, "isoformat"):
        return str(value.isoformat())
    return repr(value)


def _digest_payload(value: Any) -> str:
    encoded = json.dumps(
        value,
        sort_keys=True,
        ensure_ascii=True,
        separators=(",", ":"),
        default=_json_default,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def digest_tables(tables: dict[DatasetKey, pl.DataFrame]) -> str:
    digest = hashlib.sha256()
    for key in sorted(tables):
        frame = strip_internal_columns(tables[key])
        digest.update(f"{key.source_id}|{key.sheet}|{frame.columns}".encode())
        for row in frame.iter_rows():
            digest.update(
                json.dumps(
                    row, ensure_ascii=True, default=_json_default, separators=(",", ":")
                ).encode("utf-8")
            )
    return digest.hexdigest()


def digest_sources(bindings: tuple[SourceBinding, ...]) -> str:
    return _digest_payload(
        [
            {
                "source_id": str(binding.source_id),
                "sha256": binding.fingerprint.sha256,
                "size": binding.fingerprint.size_bytes,
            }
            for binding in sorted(bindings, key=lambda item: str(item.source_id))
        ]
    )


def load_plan_tables(
    plan: OperationPlan,
    bindings: tuple[SourceBinding, ...],
    *,
    path_overrides: dict[UUID, Path] | None = None,
) -> dict[DatasetKey, pl.DataFrame]:
    """Load only plan-relevant tables, with all paths provided by trusted bindings."""
    by_id = {binding.source_id: binding for binding in bindings}
    expected_ids = {source.source_id for source in plan.source_files}
    if set(by_id) != expected_ids:
        raise InvalidPlanError("Source bindings do not exactly match the plan sources.")
    tables: dict[DatasetKey, pl.DataFrame] = {}
    for source in plan.source_files:
        binding = by_id[source.source_id]
        if binding.fingerprint.sha256 != source.sha256:
            raise SourceChangedError("The bound source hash does not match the operation plan.")
        path = (path_overrides or {}).get(source.source_id, binding.path)
        target_sheets = {
            step.target.sheet
            for step in plan.steps
            if step.enabled and step.target.source_id == source.source_id and step.target.sheet
        }
        suffix = path.suffix.casefold()
        if suffix == ".csv":
            tables[DatasetKey(source.source_id, "CSV")] = read_tabular_source(path)
            continue
        if suffix not in {".xlsx", ".xlsm"}:
            raise InvalidPlanError("A source binding has an unsupported file format.")
        if (
            any(
                step.enabled
                and step.target.source_id == source.source_id
                and step.operation == "tables.merge"
                for step in plan.steps
            )
            or not target_sheets
        ):
            workbook = openpyxl.load_workbook(path, read_only=True, keep_links=False)
            try:
                target_sheets = set(workbook.sheetnames)
            finally:
                workbook.close()
        for sheet_name in sorted(target_sheets):
            tables[DatasetKey(source.source_id, sheet_name)] = read_workbook_table(path, sheet_name)
    return tables


class PreviewEngine:
    """Generate a change preview and bind it to plan, source, and deterministic output."""

    def __init__(self, registry: OperationRegistry, *, max_records: int = 10_000) -> None:
        self.registry = registry
        self.max_records = max_records

    def generate(
        self, plan: OperationPlan, bindings: tuple[SourceBinding, ...]
    ) -> tuple[PreviewResult, PlanRunResult]:
        for binding in bindings:
            verify_fingerprint(binding.path, binding.fingerprint)
        tables = load_plan_tables(plan, bindings)
        run = PlanRunner(self.registry, max_preview_records=self.max_records).run(plan, tables)
        for binding in bindings:
            verify_fingerprint(binding.path, binding.fingerprint)
        plan_digest = hashlib.sha256(plan.model_dump_json().encode("utf-8")).hexdigest()
        source_digest = digest_sources(bindings)
        result_digest = digest_tables(run.tables)
        preview_digest = _digest_payload(
            {
                "job_id": str(plan.job_id),
                "plan": plan_digest,
                "sources": source_digest,
                "result": result_digest,
                "changes": run.total_change_count,
            }
        )
        preview = PreviewResult(
            job_id=plan.job_id,
            created_at=datetime.now(UTC),
            changes=run.changes,
            total_change_count=run.total_change_count,
            sampled=run.total_change_count > len(run.changes),
            steps=tuple(
                PreviewStepSummary(
                    step_id=step.step_id,
                    operation=step.operation,
                    change_count=step.change_count,
                    warnings=step.warnings,
                )
                for step in run.steps
            ),
            plan_digest=plan_digest,
            source_digest=source_digest,
            result_digest=result_digest,
            preview_digest=preview_digest,
        )
        return (preview, run)


def validate_approval(
    plan: OperationPlan, preview: PreviewResult, approval: ExecutionApproval
) -> None:
    if not approval.approved:
        raise InvalidPlanError("Execution requires explicit approval.")
    if approval.job_id != plan.job_id or preview.job_id != plan.job_id:
        raise InvalidPlanError("Approval does not belong to this job.")
    if approval.preview_digest != preview.preview_digest:
        raise InvalidPlanError("Approval does not match the current preview.")
    required = {
        step.step_id
        for step in plan.steps
        if step.enabled and step.destructive and step.confirmation_required
    }
    missing = sorted(required - approval.confirmed_destructive_steps)
    if missing:
        raise InvalidPlanError(f"Destructive steps require confirmation: {', '.join(missing)}")
    known_changes = {change.change_id for change in preview.changes}
    unknown_rejections = approval.rejected_change_ids - known_changes
    if unknown_rejections:
        raise InvalidPlanError("Approval rejects unknown or unsampled preview changes.")
    rejected = [
        change for change in preview.changes if change.change_id in approval.rejected_change_ids
    ]
    if any(change.kind != ChangeKind.CELL_CHANGED for change in rejected):
        raise InvalidPlanError("Individual rejection is supported only for value changes.")


def apply_cell_rejections(
    run: PlanRunResult,
    preview: PreviewResult,
    approval: ExecutionApproval,
) -> dict[DatasetKey, pl.DataFrame]:
    """Restore explicitly rejected cell changes by stable internal row ID."""
    tables = dict(run.tables)
    # Undo later transformations first so rejecting several sequential edits to the
    # same cell restores the value that existed before the earliest rejected edit.
    for change in reversed(preview.changes):
        if change.change_id not in approval.rejected_change_ids:
            continue
        if (
            change.kind != ChangeKind.CELL_CHANGED
            or change.column is None
            or change.internal_row_id is None
        ):
            raise InvalidPlanError("This change cannot be rejected individually.")
        key = DatasetKey(change.source_id, change.sheet)
        frame = tables[key]
        if ROW_ID_COLUMN not in frame.columns:
            raise InvalidPlanError("Stable row metadata is missing during approval application.")
        frame = frame.with_columns(
            pl.when(pl.col(ROW_ID_COLUMN) == change.internal_row_id)
            .then(pl.lit(change.original_value))
            .otherwise(pl.col(change.column))
            .alias(change.column)
        )
        tables[key] = frame
    return tables
