"""Workspace-confined CSV and workbook split/merge workflows."""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import Any, cast

import openpyxl
import polars as pl
from openpyxl.cell.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from sheetpilot.app.config import SecurityLimits
from sheetpilot.core.exceptions import InvalidPlanError, OutputCollisionError
from sheetpilot.engines.csv_engine import read_csv, write_safe_csv
from sheetpilot.engines.openpyxl_export import modify_workbook_copy
from sheetpilot.engines.tabular_io import read_workbook_table
from sheetpilot.engines.xlsxwriter_engine import write_new_workbook
from sheetpilot.operations.merging import (
    MergeTablesOperation,
    MergeTablesParameters,
    SchemaMode,
)
from sheetpilot.operations.splitting import SplitByCategoryOperation, SplitByCategoryParameters
from sheetpilot.security.archive_guard import inspect_ooxml_archive
from sheetpilot.security.path_guard import ensure_within, sanitize_filename


def _copy_worksheet(source: Worksheet, target: Worksheet) -> None:
    target.sheet_format = copy(source.sheet_format)
    target.sheet_properties = copy(source.sheet_properties)
    target.page_margins = copy(source.page_margins)
    target.page_setup = copy(source.page_setup)
    target.freeze_panes = source.freeze_panes
    target.protection = copy(source.protection)
    for row in source.iter_rows():
        for source_cell in row:
            if not isinstance(source_cell, Cell):
                continue
            real_source = source_cell
            target_cell = cast(
                Cell,
                target.cell(row=int(real_source.row), column=int(real_source.column)),
            )
            target_cell.value = real_source.value
            if real_source.has_style:
                source_any = cast(Any, real_source)
                target_any = cast(Any, target_cell)
                target_any.font = copy(source_any.font)
                target_any.fill = copy(source_any.fill)
                target_any.border = copy(source_any.border)
                target_any.alignment = copy(source_any.alignment)
                target_any.number_format = source_any.number_format
                target_any.protection = copy(source_any.protection)
            if real_source.comment:
                target_cell.comment = copy(real_source.comment)
            if real_source.hyperlink:
                cast(Any, target_cell)._hyperlink = copy(real_source.hyperlink)
    for column_key, column_dimension in source.column_dimensions.items():
        target.column_dimensions[column_key] = copy(column_dimension)
    for row_key, row_dimension in source.row_dimensions.items():
        target.row_dimensions[row_key] = copy(row_dimension)
    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))


class SpreadsheetWorkflowService:
    """Run format workflows only on paths inside an owned temporary workspace."""

    def __init__(self, workspace_root: Path, limits: SecurityLimits | None = None) -> None:
        self.workspace_root = workspace_root.resolve()
        self.limits = limits or SecurityLimits()

    def _path(self, path: Path) -> Path:
        return ensure_within(path, self.workspace_root)

    def merge_csv_files(
        self,
        sources: list[Path],
        destination: Path,
        *,
        schema_mode: str = "strict",
        source_column: str | None = None,
    ) -> Path:
        tables = {path.name: read_csv(self._path(path)) for path in sources}
        parameters = MergeTablesParameters(
            schema_mode=SchemaMode(schema_mode), source_column=source_column
        )
        merged = MergeTablesOperation().execute(tables, parameters).frame
        return write_safe_csv(merged, self._path(destination))

    def combine_selected_worksheets(
        self,
        source: Path,
        sheet_names: list[str],
        *,
        schema_mode: str = "strict",
        source_column: str | None = None,
    ) -> pl.DataFrame:
        guarded = self._path(source)
        tables = {name: read_workbook_table(guarded, name) for name in sheet_names}
        return (
            MergeTablesOperation()
            .execute(
                tables,
                MergeTablesParameters(
                    schema_mode=SchemaMode(schema_mode), source_column=source_column
                ),
            )
            .frame
        )

    def merge_excel_files(
        self,
        selections: dict[Path, list[str]],
        destination: Path,
        *,
        schema_mode: str = "strict",
        source_column: str | None = None,
    ) -> Path:
        tables: dict[str, pl.DataFrame] = {}
        for path, sheets in selections.items():
            guarded = self._path(path)
            for sheet in sheets:
                tables[f"{path.name}:{sheet}"] = read_workbook_table(guarded, sheet)
        merged = (
            MergeTablesOperation()
            .execute(
                tables,
                MergeTablesParameters(
                    schema_mode=SchemaMode(schema_mode), source_column=source_column
                ),
            )
            .frame
        )
        return write_new_workbook({"Merged Data": merged}, self._path(destination))

    def split_table_to_workbook(
        self,
        frame: pl.DataFrame,
        destination: Path,
        parameters: SplitByCategoryParameters,
    ) -> Path:
        result = SplitByCategoryOperation().execute(frame, parameters)
        return write_new_workbook(result.auxiliary_tables, self._path(destination))

    def csv_to_excel(self, source: Path, destination: Path) -> Path:
        frame = read_csv(self._path(source))
        return write_new_workbook({"Data": frame}, self._path(destination))

    def create_clean_workbook(self, tables: dict[str, pl.DataFrame], destination: Path) -> Path:
        return write_new_workbook(tables, self._path(destination))

    def rename_sheets_on_copy(
        self, source: Path, destination: Path, renames: dict[str, str]
    ) -> Path:
        return modify_workbook_copy(
            self._path(source), self._path(destination), replacements={}, sheet_renames=renames
        )

    def add_summary_sheet_on_copy(
        self,
        source: Path,
        destination: Path,
        summary_name: str,
        summary: pl.DataFrame,
    ) -> Path:
        guarded = self._path(source)
        workbook = openpyxl.load_workbook(guarded, read_only=True, keep_links=False)
        try:
            if summary_name in workbook.sheetnames:
                raise InvalidPlanError("The summary sheet name already exists.")
        finally:
            workbook.close()
        return modify_workbook_copy(
            guarded,
            self._path(destination),
            replacements={summary_name: summary},
        )

    def export_selected_sheets(
        self, source: Path, destination: Path, selected_sheets: list[str]
    ) -> Path:
        source = self._path(source)
        destination = self._path(destination)
        if destination.exists():
            raise OutputCollisionError("Selected-sheet export cannot overwrite a file.")
        inspection = inspect_ooxml_archive(source, self.limits)
        if inspection.macro_present:
            raise InvalidPlanError("Selected-sheet export is blocked for macro-enabled workbooks.")
        source_book = openpyxl.load_workbook(source, data_only=False, keep_links=False)
        output_book = openpyxl.Workbook()
        active_sheet = output_book.active
        if active_sheet is not None:
            output_book.remove(active_sheet)
        try:
            missing = sorted(set(selected_sheets) - set(source_book.sheetnames))
            if missing:
                raise InvalidPlanError(f"Sheets to export are missing: {', '.join(missing)}")
            for name in selected_sheets:
                target = output_book.create_sheet(name)
                _copy_worksheet(source_book[name], target)
            output_book.save(destination)
        except BaseException:
            output_book.close()
            source_book.close()
            destination.unlink(missing_ok=True)
            raise
        output_book.close()
        source_book.close()
        return destination

    def split_workbook(
        self, source: Path, output_directory: Path, selected_sheets: list[str] | None = None
    ) -> list[Path]:
        source = self._path(source)
        output_directory = self._path(output_directory)
        output_directory.mkdir(parents=True, exist_ok=True)
        inspection = inspect_ooxml_archive(source, self.limits)
        if inspection.macro_present:
            raise InvalidPlanError("Workbook splitting is blocked for macro-enabled workbooks.")
        workbook = openpyxl.load_workbook(source, data_only=False, keep_links=False)
        names = selected_sheets or workbook.sheetnames
        workbook.close()
        outputs: list[Path] = []
        for name in names:
            filename = f"{sanitize_filename(source.stem)}_{sanitize_filename(name)}.xlsx"
            destination = self._path(output_directory / filename)
            outputs.append(self.export_selected_sheets(source, destination, [name]))
        return outputs
