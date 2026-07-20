from __future__ import annotations

from pathlib import Path
from uuid import uuid4

import openpyxl
import polars as pl
import pytest
from openpyxl.styles import PatternFill

from sheetpilot.app.config import AppConfig
from sheetpilot.core.exceptions import (
    FormulaPreservationRiskError,
    WorkbookFeaturePreservationRiskError,
)
from sheetpilot.core.executor import JobExecutor
from sheetpilot.core.plan_schema import (
    OperationPlan,
    OutputSettings,
    PlanStep,
    SourceReference,
    StepTarget,
)
from sheetpilot.core.preview_engine import ExecutionApproval, PreviewEngine, SourceBinding
from sheetpilot.engines.openpyxl_export import modify_workbook_copy
from sheetpilot.operations.registry import build_default_registry
from sheetpilot.security.hashing import fingerprint_file


def _config(tmp_path: Path) -> AppConfig:
    data = tmp_path / "data"
    return AppConfig(
        data_dir=data,
        backup_dir=data / "backups",
        temp_dir=data / "temp",
        database_path=data / "metadata.sqlite3",
        log_dir=data / "logs",
    )


def test_text_cleanup_preserves_unchanged_existing_xlsx_formula(tmp_path: Path) -> None:
    source = tmp_path / "formula-source.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Name", "Qty", "Price", "Total"])
    worksheet.append([" Ada ", 2, 5, "=B2*C2"])
    worksheet["D2"].fill = PatternFill("solid", fgColor="FFFFFF00")
    workbook.save(source)
    workbook.close()
    source_before = source.read_bytes()

    source_id = uuid4()
    fingerprint = fingerprint_file(source)
    plan = OperationPlan(
        job_name="Clean names without damaging formulas",
        source_files=[
            SourceReference(
                source_id=source_id,
                file_name=source.name,
                sha256=fingerprint.sha256,
                sheet_names=["Data"],
            )
        ],
        steps=[
            PlanStep(
                step_id="clean-name",
                operation="text.clean",
                parameters={"columns": ["Name"], "actions": [{"kind": "trim"}]},
                target=StepTarget(source_id=source_id, sheet="Data", columns=["Name"]),
                explanation="Trim surrounding whitespace from names",
            )
        ],
        output=OutputSettings(
            output_name="cleaned",
            format="xlsx",
            preserve_formatting=True,
        ),
    )
    binding = SourceBinding(source_id=source_id, path=source, fingerprint=fingerprint)
    registry = build_default_registry()
    preview, _ = PreviewEngine(registry).generate(plan, (binding,))

    result = JobExecutor(_config(tmp_path), registry).execute(
        plan,
        (binding,),
        preview,
        ExecutionApproval(
            job_id=plan.job_id,
            preview_digest=preview.preview_digest,
            approved=True,
        ),
        tmp_path / "output",
    )

    output = openpyxl.load_workbook(result.output.path, data_only=False)
    try:
        assert output["Data"]["A2"].value == "Ada"
        assert output["Data"]["D2"].data_type == "f"
        assert output["Data"]["D2"].value == "=B2*C2"
        assert output["Data"]["D2"].fill.fgColor.rgb == "FFFFFF00"
    finally:
        output.close()
    assert source.read_bytes() == source_before


def test_existing_formula_movement_fails_closed(tmp_path: Path) -> None:
    source = tmp_path / "formula-source.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Name", "Qty", "Total"])
    worksheet.append(["Ada", 2, "=B2*2"])
    worksheet.append(["Grace", 3, "=B3*2"])
    workbook.save(source)
    workbook.close()
    source_before = source.read_bytes()
    destination = tmp_path / "unsafe-result.xlsx"

    with pytest.raises(
        FormulaPreservationRiskError,
        match="cannot safely rebase existing formulas",
    ):
        modify_workbook_copy(
            source,
            destination,
            replacements={
                "Data": pl.DataFrame(
                    {
                        "Name": ["Grace", "Ada"],
                        "Qty": [3, 2],
                        "Total": ["=B3*2", "=B2*2"],
                    }
                )
            },
        )

    assert not destination.exists()
    assert source.read_bytes() == source_before


def test_merged_sheet_replacement_fails_with_typed_preservation_error(tmp_path: Path) -> None:
    source = tmp_path / "merged-source.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["Customer", "Region"])
    worksheet.append(["Ada", "East"])
    worksheet.merge_cells("A2:B2")
    workbook.save(source)
    workbook.close()
    source_before = source.read_bytes()
    destination = tmp_path / "unsafe-merged-result.xlsx"

    with pytest.raises(
        WorkbookFeaturePreservationRiskError,
        match="cannot safely replace that table",
    ):
        modify_workbook_copy(
            source,
            destination,
            replacements={"Data": pl.DataFrame({"Customer": ["Ada"], "Region": ["East"]})},
        )

    assert not destination.exists()
    assert source.read_bytes() == source_before
