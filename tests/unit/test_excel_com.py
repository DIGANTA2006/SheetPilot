from __future__ import annotations

import shutil
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

import openpyxl
import pytest

import sheetpilot.integrations.excel_com as excel_com_module
from sheetpilot.core.exceptions import (
    ExcelComOperationError,
    ExcelComUnavailableError,
    OutputCollisionError,
    PathSecurityError,
    TrustedMacroError,
)
from sheetpilot.integrations.excel_com import ExcelComEngine, TrustedMacroApproval
from sheetpilot.security.hashing import fingerprint_file

_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_VBA_CONTENT_TYPE = "application/vnd.ms-office.vbaProject"
_VBA_RELATIONSHIP = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"
_MACRO_WORKBOOK_TYPE = "application/vnd.ms-excel.sheet.macroEnabled.main+xml"
_VALID_PDF = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\n%%EOF\n"


def _create_workbook(path: Path) -> None:
    base = path if path.suffix.casefold() == ".xlsx" else path.with_suffix(".xlsx")
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["ID", "Name"])
    worksheet.append([1, "Ada"])
    workbook.save(base)
    workbook.close()
    if path.suffix.casefold() != ".xlsm":
        return

    with ZipFile(base) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}
    content_types = ElementTree.fromstring(members["[Content_Types].xml"])
    for node in content_types:
        if node.attrib.get("PartName") == "/xl/workbook.xml":
            node.set("ContentType", _MACRO_WORKBOOK_TYPE)
    ElementTree.SubElement(
        content_types,
        f"{{{_CONTENT_TYPES_NS}}}Override",
        PartName="/xl/vbaProject.bin",
        ContentType=_VBA_CONTENT_TYPE,
    )
    members["[Content_Types].xml"] = ElementTree.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )
    relationships = ElementTree.fromstring(members["xl/_rels/workbook.xml.rels"])
    ElementTree.SubElement(
        relationships,
        f"{{{_RELATIONSHIPS_NS}}}Relationship",
        Id="rIdSheetPilotVba",
        Type=_VBA_RELATIONSHIP,
        Target="vbaProject.bin",
    )
    members["xl/_rels/workbook.xml.rels"] = ElementTree.tostring(
        relationships, encoding="utf-8", xml_declaration=True
    )
    members["xl/vbaProject.bin"] = b"Trusted test macro payload; never executed by the fixture."
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)
    base.unlink()


class FakePivotCache:
    def __init__(self) -> None:
        self.refreshed = False

    def Refresh(self) -> None:  # noqa: N802
        self.refreshed = True


class FakePivot:
    def __init__(self) -> None:
        self.cache = FakePivotCache()

    def PivotCache(self) -> FakePivotCache:  # noqa: N802
        return self.cache


class FakeWorksheet:
    def __init__(self) -> None:
        self.pivots: dict[str, FakePivot] = {}

    def PivotTables(self, name: str) -> FakePivot:  # noqa: N802
        return self.pivots.setdefault(name, FakePivot())

    def ExportAsFixedFormat(self, format_code: int, destination: str) -> None:  # noqa: N802
        assert format_code == 0
        Path(destination).write_bytes(_VALID_PDF)


class FakeWorkbook:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.Name = path.name
        self.sheets: dict[str, FakeWorksheet] = {}
        self.closed = False
        self.saved_destinations: list[Path] = []

    def SaveCopyAs(self, destination: str) -> None:  # noqa: N802
        self.saved_destinations.append(Path(destination))
        shutil.copy2(self.path, destination)

    def Worksheets(self, name: str) -> FakeWorksheet:  # noqa: N802
        return self.sheets.setdefault(name, FakeWorksheet())

    def Close(self, **kwargs: Any) -> None:  # noqa: N802
        assert not kwargs["SaveChanges"]
        self.closed = True


class FakeWorkbooks:
    def __init__(self) -> None:
        self.opened: FakeWorkbook | None = None

    def Open(self, path: str, **kwargs: Any) -> FakeWorkbook:  # noqa: N802
        assert kwargs["UpdateLinks"] == 0
        assert not kwargs["ReadOnly"]
        self.opened = FakeWorkbook(Path(path))
        return self.opened


class FakeExcel:
    def __init__(self) -> None:
        self.Workbooks = FakeWorkbooks()
        self.AutomationSecurity = 0
        self.Visible = True
        self.DisplayAlerts = True
        self.EnableEvents = True
        self.recalculated = False
        self.runs: list[str] = []
        self.quit_called = False

    def CalculateFullRebuild(self) -> None:  # noqa: N802
        self.recalculated = True

    def Run(self, macro: str) -> None:  # noqa: N802
        self.runs.append(macro)

    def Quit(self) -> None:  # noqa: N802
        self.quit_called = True


class InvalidPdfWorksheet(FakeWorksheet):
    def ExportAsFixedFormat(self, format_code: int, destination: str) -> None:  # noqa: N802
        assert format_code == 0
        Path(destination).write_bytes(b"not a pdf")


class InvalidPdfWorkbook(FakeWorkbook):
    def Worksheets(self, name: str) -> FakeWorksheet:  # noqa: N802
        return self.sheets.setdefault(name, InvalidPdfWorksheet())


class InvalidPdfWorkbooks(FakeWorkbooks):
    def Open(self, path: str, **kwargs: Any) -> FakeWorkbook:  # noqa: N802
        assert kwargs["UpdateLinks"] == 0
        self.opened = InvalidPdfWorkbook(Path(path))
        return self.opened


class InvalidPdfExcel(FakeExcel):
    def __init__(self) -> None:
        super().__init__()
        self.Workbooks = InvalidPdfWorkbooks()


class CleanupFailureExcel(FakeExcel):
    def __init__(self, *, fail_operation: bool = False) -> None:
        self._automation_security = 0
        self.fail_cleanup = False
        self.fail_operation = fail_operation
        super().__init__()

    @property
    def AutomationSecurity(self) -> int:  # noqa: N802
        return self._automation_security

    @AutomationSecurity.setter
    def AutomationSecurity(self, value: int) -> None:  # noqa: N802
        if self.fail_cleanup and value == 3:
            raise RuntimeError("cleanup reset failed")
        self._automation_security = value

    def CalculateFullRebuild(self) -> None:  # noqa: N802
        self.fail_cleanup = True
        if self.fail_operation:
            raise RuntimeError("primary recalculation failed")
        self.recalculated = True


def _workbook(tmp_path: Path, suffix: str = ".xlsx") -> tuple[Path, Path, Path]:
    workspace = tmp_path / "workspace"
    output = tmp_path / "output"
    workspace.mkdir(parents=True)
    output.mkdir(parents=True)
    source = workspace / f"working{suffix}"
    _create_workbook(source)
    return source, workspace, output


def test_recalculation_uses_disabled_macros_and_distinct_copy(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    application = FakeExcel()
    engine = ExcelComEngine(lambda: application)

    artifact = engine.recalculate_to_copy(
        source,
        output / "recalculated.xlsx",
        workspace_root=workspace,
        approved_output_root=output,
    )

    assert application.recalculated
    assert application.AutomationSecurity == 3
    assert application.quit_called
    assert artifact.path.read_bytes() == source.read_bytes()
    workbook = application.Workbooks.opened
    assert workbook is not None
    assert len(workbook.saved_destinations) == 1
    staged = workbook.saved_destinations[0]
    assert staged.parent == output
    assert staged.suffix == ".xlsx"
    assert staged != artifact.path
    assert not staged.exists()


def test_pivot_refresh_is_limited_to_explicit_names(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    application = FakeExcel()
    engine = ExcelComEngine(lambda: application)

    engine.refresh_approved_pivots_to_copy(
        source,
        output / "refreshed.xlsx",
        ("Summary!SalesPivot",),
        workspace_root=workspace,
        approved_output_root=output,
    )

    workbook = application.Workbooks.opened
    assert workbook is not None
    assert workbook.sheets["Summary"].pivots["SalesPivot"].cache.refreshed


def test_trusted_macro_is_hash_bound_and_opened_with_macros_disabled(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path, ".xlsm")
    application = FakeExcel()
    engine = ExcelComEngine(lambda: application)
    fingerprint = fingerprint_file(source)
    approval = TrustedMacroApproval(
        workbook_sha256=fingerprint.sha256,
        macro_name="ApprovedModule.Clean",
        explicitly_confirmed=True,
    )

    engine.run_trusted_macro_to_copy(
        source,
        output / "macro-result.xlsm",
        approval,
        workspace_root=workspace,
        approved_output_root=output,
    )

    assert application.runs == ["'working.xlsm'!ApprovedModule.Clean"]
    assert application.AutomationSecurity == 3


def test_trusted_macro_rejects_a_different_workbook_hash(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path, ".xlsm")
    engine = ExcelComEngine(FakeExcel)
    approval = TrustedMacroApproval(
        workbook_sha256="0" * 64,
        macro_name="ApprovedModule.Clean",
        explicitly_confirmed=True,
    )

    with pytest.raises(TrustedMacroError):
        engine.run_trusted_macro_to_copy(
            source,
            output / "macro-result.xlsm",
            approval,
            workspace_root=workspace,
            approved_output_root=output,
        )


def test_com_outputs_never_overwrite_existing_files(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    destination = output / "existing.xlsx"
    destination.write_bytes(b"existing")
    engine = ExcelComEngine(FakeExcel)

    with pytest.raises(OutputCollisionError):
        engine.recalculate_to_copy(
            source,
            destination,
            workspace_root=workspace,
            approved_output_root=output,
        )

    assert destination.read_bytes() == b"existing"


def test_pdf_export_is_staged_and_signature_validated(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    application = FakeExcel()
    engine = ExcelComEngine(lambda: application)

    artifact = engine.export_sheet_pdf(
        source,
        output / "Data.pdf",
        "Data",
        workspace_root=workspace,
        approved_output_root=output,
    )

    assert artifact.path.read_bytes() == _VALID_PDF
    assert artifact.fingerprint == fingerprint_file(artifact.path)
    assert not tuple(output.glob(".*.partial.pdf"))


def test_invalid_pdf_is_not_published(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    destination = output / "Data.pdf"

    with pytest.raises(ExcelComOperationError, match=r"truncated PDF|invalid PDF"):
        ExcelComEngine(InvalidPdfExcel).export_sheet_pdf(
            source,
            destination,
            "Data",
            workspace_root=workspace,
            approved_output_root=output,
        )

    assert not destination.exists()
    assert not tuple(output.glob(".*.partial.pdf"))


def test_cleanup_failure_attempts_close_and_quit_without_publishing(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    application = CleanupFailureExcel()
    destination = output / "recalculated.xlsx"

    with pytest.raises(ExcelComOperationError, match="did not close cleanly"):
        ExcelComEngine(lambda: application).recalculate_to_copy(
            source,
            destination,
            workspace_root=workspace,
            approved_output_root=output,
        )

    workbook = application.Workbooks.opened
    assert workbook is not None and workbook.closed
    assert application.quit_called
    assert not destination.exists()


def test_cleanup_failure_does_not_mask_primary_com_failure(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    application = CleanupFailureExcel(fail_operation=True)

    with pytest.raises(ExcelComOperationError) as raised:
        ExcelComEngine(lambda: application).recalculate_to_copy(
            source,
            output / "recalculated.xlsx",
            workspace_root=workspace,
            approved_output_root=output,
        )

    assert isinstance(raised.value.__cause__, RuntimeError)
    assert str(raised.value.__cause__) == "primary recalculation failed"
    workbook = application.Workbooks.opened
    assert workbook is not None and workbook.closed
    assert application.quit_called


def test_unavailable_excel_has_precise_typed_failure(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)

    def unavailable() -> Any:
        raise RuntimeError("class not registered")

    engine = ExcelComEngine(unavailable)
    assert not engine.available
    with pytest.raises(ExcelComUnavailableError, match="could not be started"):
        engine.recalculate_to_copy(
            source,
            output / "recalculated.xlsx",
            workspace_root=workspace,
            approved_output_root=output,
        )


def test_real_adapter_initializes_and_releases_com_apartment(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    events: list[str] = []
    application = FakeExcel()
    original_quit = application.Quit

    def quit_application() -> None:
        events.append("quit")
        original_quit()

    application.Quit = quit_application
    pythoncom = SimpleNamespace(
        CoInitialize=lambda: events.append("initialize"),
        CoUninitialize=lambda: events.append("uninitialize"),
    )
    client = SimpleNamespace(
        DispatchEx=lambda name: events.append(f"dispatch:{name}") or application
    )

    def import_module(name: str) -> object:
        if name == "pythoncom":
            return pythoncom
        if name == "win32com.client":
            return client
        raise ImportError(name)

    monkeypatch.setattr(excel_com_module.importlib, "import_module", import_module)

    assert ExcelComEngine().available
    assert events == [
        "initialize",
        "dispatch:Excel.Application",
        "quit",
        "uninitialize",
    ]


def test_output_path_traversal_and_wrong_extensions_are_rejected(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    engine = ExcelComEngine(FakeExcel)

    with pytest.raises(PathSecurityError):
        engine.recalculate_to_copy(
            source,
            tmp_path / "escaped.xlsx",
            workspace_root=workspace,
            approved_output_root=output,
        )
    with pytest.raises(PathSecurityError):
        engine.recalculate_to_copy(
            source,
            output / "wrong.xlsm",
            workspace_root=workspace,
            approved_output_root=output,
        )


def test_trusted_macro_requires_macro_content_and_matching_extension(tmp_path: Path) -> None:
    source, workspace, output = _workbook(tmp_path)
    approval = TrustedMacroApproval(
        workbook_sha256=fingerprint_file(source).sha256,
        macro_name="ApprovedModule.Clean",
        explicitly_confirmed=True,
    )
    with pytest.raises(TrustedMacroError, match="macro-enabled"):
        ExcelComEngine(FakeExcel).run_trusted_macro_to_copy(
            source,
            output / "macro-result.xlsm",
            approval,
            workspace_root=workspace,
            approved_output_root=output,
        )

    macro_source, macro_workspace, macro_output = _workbook(tmp_path / "macro", ".xlsm")
    macro_approval = approval.model_copy(
        update={"workbook_sha256": fingerprint_file(macro_source).sha256}
    )
    with pytest.raises(PathSecurityError):
        ExcelComEngine(FakeExcel).run_trusted_macro_to_copy(
            macro_source,
            macro_output / "macro-result.xlsx",
            macro_approval,
            workspace_root=macro_workspace,
            approved_output_root=macro_output,
        )


def test_trusted_macro_escapes_workbook_name_in_qualified_target(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    output = tmp_path / "output"
    workspace.mkdir(parents=True)
    output.mkdir(parents=True)
    source = workspace / "working'copy.xlsm"
    _create_workbook(source)
    application = FakeExcel()
    approval = TrustedMacroApproval(
        workbook_sha256=fingerprint_file(source).sha256,
        macro_name="ApprovedModule.Clean",
        explicitly_confirmed=True,
    )

    ExcelComEngine(lambda: application).run_trusted_macro_to_copy(
        source,
        output / "macro-result.xlsm",
        approval,
        workspace_root=workspace,
        approved_output_root=output,
    )

    assert application.runs == ["'working''copy.xlsm'!ApprovedModule.Clean"]
