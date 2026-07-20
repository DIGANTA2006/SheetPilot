from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import polars as pl
import pytest
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
from pydantic import ValidationError

from sheetpilot.core.exceptions import InvalidPlanError
from sheetpilot.core.file_workflows import SpreadsheetWorkflowService
from sheetpilot.engines.csv_engine import write_safe_csv
from sheetpilot.engines.duckdb_engine import DuckDBEngine
from sheetpilot.engines.formulas import GeneratedFormulaSpec
from sheetpilot.engines.openpyxl_export import (
    ColumnFormatSpec,
    NumberFormat,
    modify_workbook_copy,
)
from sheetpilot.engines.pandas_engine import from_pandas, to_pandas
from sheetpilot.engines.tabular_io import validate_output_file
from sheetpilot.engines.xlsxwriter_engine import write_new_workbook
from sheetpilot.operations.merging import MergeTablesOperation, MergeTablesParameters
from sheetpilot.operations.splitting import SplitByCategoryOperation, SplitByCategoryParameters
from sheetpilot.security.hashing import fingerprint_file


def test_safe_csv_and_polished_xlsx_exports_never_execute_text(tmp_path: Path) -> None:
    frame = pl.DataFrame({"=Payload": ["=2+2"], "Qty": [2], "Price": [5]})
    csv_path = write_safe_csv(frame, tmp_path / "safe.csv")
    with csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    assert rows[0][0] == "'=Payload"
    assert rows[1][0] == "'=2+2"

    xlsx_path = write_new_workbook(
        {"Data": frame},
        tmp_path / "safe.xlsx",
        formula_specs={
            "Data": [
                GeneratedFormulaSpec(
                    output_column="Total",
                    kind="quantity_price",
                    operand_columns=["Qty", "Price"],
                    number_format="0.00",
                )
            ]
        },
    )
    validate_output_file(xlsx_path)
    workbook = openpyxl.load_workbook(xlsx_path, data_only=False)
    try:
        sheet = workbook["Data"]
        assert sheet["A2"].data_type == "s"
        assert sheet["A2"].value == "=2+2"
        assert sheet["D2"].data_type == "f"
        assert sheet["D2"].value == "=B2*C2"
        assert sheet.freeze_panes == "A2"
        assert sheet["A1"].fill.fgColor.rgb == "FF1D4ED8"
    finally:
        workbook.close()


def test_exports_reject_case_insensitive_column_collisions(tmp_path: Path) -> None:
    destination = tmp_path / "unsafe.csv"
    with pytest.raises(InvalidPlanError, match="unique without regard to letter case"):
        write_safe_csv(pl.DataFrame({"Name": ["Ada"], "name": ["Grace"]}), destination)
    assert not destination.exists()


def test_structured_formula_schema_rejects_raw_formula() -> None:
    try:
        GeneratedFormulaSpec.model_validate(
            {
                "output_column": "Unsafe",
                "kind": "addition",
                "operand_columns": ["A", "B"],
                "formula": "=EXEC()",
            }
        )
    except ValidationError:
        pass
    else:
        raise AssertionError("raw formula field was accepted")


def test_copy_only_openpyxl_edit_preserves_source_and_workbook_features(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["A", "Value"])
    data.append([1, 2])
    data["A1"].fill = PatternFill("solid", fgColor="FFCC0000")
    data["A2"].comment = Comment("Synthetic fixture comment", "SheetPilot")
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    workbook.save(source)
    workbook.close()
    before = fingerprint_file(source)

    destination = tmp_path / "modified.xlsx"
    modify_workbook_copy(
        source,
        destination,
        replacements={"Data": pl.DataFrame({"A": [3], "Value": [4]})},
        formula_specs={
            "Data": [
                GeneratedFormulaSpec(
                    output_column="Sum", kind="addition", operand_columns=["A", "Value"]
                )
            ]
        },
        formats={
            "Data": [ColumnFormatSpec(column="Value", number_format=NumberFormat.CURRENCY_INR)]
        },
    )
    assert fingerprint_file(source) == before
    validate_output_file(destination)
    result = openpyxl.load_workbook(destination, data_only=False)
    try:
        assert result["Hidden"].sheet_state == "hidden"
        assert result["Data"]["A1"].fill.fgColor.rgb == "FFCC0000"
        assert result["Data"]["C2"].value == "=A2+B2"
        assert result["Data"]["B2"].number_format == "₹#,##0.00"
    finally:
        result.close()


def test_split_and_merge_table_operations() -> None:
    tables = {
        "January": pl.DataFrame({"ID": [1], "Value": [10]}),
        "February": pl.DataFrame({"ID": [2], "Other": [20]}),
    }
    union = MergeTablesOperation().execute(
        tables,
        MergeTablesParameters(schema_mode="union", source_column="Source"),
    )
    assert union.frame.height == 2
    assert set(union.frame.columns) == {"ID", "Value", "Other", "Source"}

    split = SplitByCategoryOperation().execute(
        pl.DataFrame({"Region": ["North", "South", "North"], "Value": [1, 2, 3]}),
        SplitByCategoryParameters(category_column="Region", table_prefix="Region - "),
    )
    assert set(split.auxiliary_tables) == {"Region - North", "Region - South"}
    assert split.auxiliary_tables["Region - North"].height == 2


def test_workspace_confined_file_workflows(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    first_csv = workspace / "one.csv"
    second_csv = workspace / "two.csv"
    first_csv.write_text("ID,Name\n1,Ada\n", encoding="utf-8")
    second_csv.write_text("ID,Name\n2,Grace\n", encoding="utf-8")
    service = SpreadsheetWorkflowService(workspace)
    merged_csv = service.merge_csv_files(
        [first_csv, second_csv], workspace / "merged.csv", source_column="Source"
    )
    assert pl.read_csv(merged_csv).height == 2
    converted = service.csv_to_excel(first_csv, workspace / "converted.xlsx")
    validate_output_file(converted)

    source_book = workspace / "book.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "One"
    workbook.active.append(["ID", "Value"])
    workbook.active.append([1, 10])
    second = workbook.create_sheet("Two")
    second.append(["ID", "Value"])
    second.append([2, 20])
    workbook.save(source_book)
    workbook.close()
    source_hash = fingerprint_file(source_book)

    combined = service.combine_selected_worksheets(
        source_book, ["One", "Two"], source_column="Sheet"
    )
    assert combined.height == 2
    renamed = service.rename_sheets_on_copy(
        source_book, workspace / "renamed.xlsx", {"One": "Primary"}
    )
    renamed_book = openpyxl.load_workbook(renamed, read_only=True)
    try:
        assert "Primary" in renamed_book.sheetnames
    finally:
        renamed_book.close()
    summary = service.add_summary_sheet_on_copy(
        source_book,
        workspace / "summary.xlsx",
        "Summary",
        pl.DataFrame({"Metric": ["Rows"], "Value": [2]}),
    )
    summary_book = openpyxl.load_workbook(summary, read_only=True)
    try:
        assert "Summary" in summary_book.sheetnames
    finally:
        summary_book.close()
    selected = service.export_selected_sheets(source_book, workspace / "selected.xlsx", ["Two"])
    selected_book = openpyxl.load_workbook(selected, read_only=True)
    try:
        assert selected_book.sheetnames == ["Two"]
    finally:
        selected_book.close()
    outputs = service.split_workbook(source_book, workspace / "split")
    assert len(outputs) == 2 and all(path.exists() for path in outputs)
    assert fingerprint_file(source_book) == source_hash


def test_duckdb_large_csv_union_and_pandas_compatibility(tmp_path: Path) -> None:
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    first = workspace / "first.csv"
    second = workspace / "second.csv"
    first.write_text("ID\n1\n", encoding="utf-8")
    second.write_text("ID\n2\n", encoding="utf-8")
    result = DuckDBEngine(workspace).union_csv_files([first, second])
    assert result.get_column("ID").to_list() == [1, 2]
    pandas_frame = to_pandas(result)
    assert from_pandas(pandas_frame).equals(result)
