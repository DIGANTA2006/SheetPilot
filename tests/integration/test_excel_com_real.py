from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

from sheetpilot.core.exceptions import OutputCollisionError
from sheetpilot.integrations.excel_com import ExcelComEngine
from sheetpilot.security.hashing import FileFingerprint, fingerprint_file

pytestmark = pytest.mark.excel_com


def _real_excel_engine() -> ExcelComEngine:
    if sys.platform != "win32":
        pytest.skip("real Excel COM integration requires Windows")
    pytest.importorskip(
        "win32com.client",
        reason="real Excel COM integration requires the optional pywin32 dependency",
    )
    engine = ExcelComEngine()
    if not engine.available:
        pytest.skip("real Excel COM integration requires an installed Microsoft Excel instance")
    return engine


def _create_synthetic_working_copy(path: Path) -> FileFingerprint:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet.append(["First value", "Second value", "Calculated total"])
    worksheet.append([2, 3, "=SUM(A2:B2)"])
    worksheet.print_area = "A1:C2"
    worksheet.freeze_panes = "A2"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)
    workbook.close()
    return fingerprint_file(path)


def test_real_excel_recalculation_and_pdf_export_preserve_working_copy(
    tmp_path: Path,
) -> None:
    """Exercise real Excel only against a disposable, job-owned workbook copy."""
    engine = _real_excel_engine()
    workspace = tmp_path / "workspace"
    output_root = tmp_path / "output"
    workspace.mkdir()
    output_root.mkdir()
    working_copy = workspace / "synthetic-working-copy.xlsx"
    source_fingerprint = _create_synthetic_working_copy(working_copy)
    source_mtime_ns = working_copy.stat().st_mtime_ns

    recalculated = engine.recalculate_to_copy(
        working_copy,
        output_root / "recalculated.xlsx",
        workspace_root=workspace,
        approved_output_root=output_root,
    )

    assert recalculated.path == output_root / "recalculated.xlsx"
    assert recalculated.fingerprint == fingerprint_file(recalculated.path)
    formula_workbook = openpyxl.load_workbook(recalculated.path, data_only=False)
    try:
        assert formula_workbook["Data"]["C2"].value == "=SUM(A2:B2)"
    finally:
        formula_workbook.close()
    value_workbook = openpyxl.load_workbook(recalculated.path, data_only=True)
    try:
        assert value_workbook["Data"]["C2"].value == 5
    finally:
        value_workbook.close()

    pdf = engine.export_sheet_pdf(
        working_copy,
        output_root / "Data.pdf",
        "Data",
        workspace_root=workspace,
        approved_output_root=output_root,
    )

    pdf_bytes = pdf.path.read_bytes()
    assert pdf.path == output_root / "Data.pdf"
    assert pdf.fingerprint == fingerprint_file(pdf.path)
    assert pdf_bytes.startswith(b"%PDF-")
    assert b"%%EOF" in pdf_bytes[-2048:]

    collision = output_root / "existing.xlsx"
    collision.write_bytes(b"existing output must not be replaced")
    with pytest.raises(OutputCollisionError):
        engine.recalculate_to_copy(
            working_copy,
            collision,
            workspace_root=workspace,
            approved_output_root=output_root,
        )
    assert collision.read_bytes() == b"existing output must not be replaced"

    assert fingerprint_file(working_copy) == source_fingerprint
    assert working_copy.stat().st_mtime_ns == source_mtime_ns
    assert not tuple(output_root.glob(".*.partial.*"))
