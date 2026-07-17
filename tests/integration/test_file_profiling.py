from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.workbook.defined_name import DefinedName

from sheetpilot.core.file_profiler import FileProfiler
from sheetpilot.security.hashing import fingerprint_file


def test_csv_profile_is_aggregate_and_non_destructive(tmp_path: Path) -> None:
    source = tmp_path / "clients.csv"
    source.write_text(
        "ID,Status,Status,Email,Mobile,DOB,Payload\n"
        "1,Active,Active,valid@example.com,9876543210,2025-01-01,ordinary\n"
        "unknown,active,Active,broken@,123,not-a-date,=2+2\n"
        "unknown,active,Active,broken@,123,not-a-date,=2+2\n",
        encoding="utf-8",
    )
    before = fingerprint_file(source)
    profile = FileProfiler().profile(source)
    after = fingerprint_file(source)

    assert before == after == profile.fingerprint
    assert profile.file_type == "csv"
    assert profile.visible_sheets == ("CSV",)
    sheet = profile.sheets[0]
    assert sheet.duplicate_headers == ("Status",)
    assert sheet.duplicate_rows == 1
    assert any(column.mixed_types for column in sheet.columns)
    assert sum(column.formula_injection_count for column in sheet.columns) == 2
    assert sum(column.suspicious_email_count for column in sheet.columns) == 2
    assert sum(column.suspicious_phone_count for column in sheet.columns) == 2
    assert sum(column.invalid_date_count for column in sheet.columns) == 2
    assert {warning.code for warning in profile.warnings} >= {
        "duplicate_headers",
        "formula_injection",
        "invalid_contact_or_date_values",
    }


def test_excel_profile_reports_workbook_features_without_modification(tmp_path: Path) -> None:
    source = tmp_path / "profile.xlsx"
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["ID", "Name", "Name", "Total", "External"])
    data.append([1, "Ada", "Ada", "=1+1", "='[other.xlsx]Sheet1'!A1"])
    data.append([1, "Ada", "Ada", "=1+1", "='[other.xlsx]Sheet1'!A1"])
    data.merge_cells("A5:B5")
    data.protection.sheet = True
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Value"])
    hidden.append(["secret metadata"])
    workbook.security.lockStructure = True
    workbook.defined_names.add(DefinedName("DataRows", attr_text="'Data'!$A$2:$A$3"))
    workbook.save(source)
    workbook.close()

    before = fingerprint_file(source)
    profile = FileProfiler().profile(source)
    after = fingerprint_file(source)

    assert before == after
    assert profile.hidden_sheets == ("Hidden",)
    assert profile.workbook_protected
    assert profile.external_links_present
    assert profile.named_ranges == ("DataRows",)
    data_profile = profile.sheets[0]
    assert data_profile.duplicate_headers == ("Name",)
    assert data_profile.duplicate_rows == 1
    assert data_profile.formula_cells == 4
    assert data_profile.merged_ranges == ("A5:B5",)
    assert data_profile.protected
    codes = {warning.code for warning in profile.warnings}
    assert {"hidden_sheet", "protected_sheet", "protected_workbook", "external_links"} <= codes
